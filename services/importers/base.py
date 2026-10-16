"""The fail-closed import pipeline.

    source spec → workbook inspection → validation report → normalized staging records
    → entity-resolution report → verification → staged ImportResult

Any `error` issue means no normalized success: `ImportResult.ok` is False and `records` is empty.
Warnings proceed only when the spec allows them. Nothing is persisted and nothing is wired into
the live M5 context or refresh; M6 sources plug in when they are introduced.
"""
from __future__ import annotations

import datetime as dt
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal, Mapping
from xml.etree.ElementTree import ParseError

from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.utils.exceptions import InvalidFileException

from .dates import resolve_business_date
from .entity_resolution import MasterIndex, resolve_values
from .excel import EMPTY_CELL, CellRead, SheetPair, WorkbookPair, external_link_targets, label_text
from .keys import keys_with_occurrence
from .models import (
    AttributeSource,
    Audience,
    BusinessDate,
    CellKind,
    Confidence,
    MatchReport,
    NormalizedRecord,
    NormalizedValue,
    Severity,
    ValidationReport,
    VerificationResult,
)
from .spec import FieldSpec, MatrixLayout, MatrixRange, MatrixSection, SourceSpec, TableLayout, cell_rc
from .validation import (
    ID_FLOAT,
    ID_MALFORMED,
    ID_PLACEHOLDER,
    TYPE_MISMATCH,
    FieldAccumulator,
    normalize_cell,
    resolve_headers,
)

SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}


@dataclass
class RuleResult:
    name: str
    checked: int = 0
    failed: int = 0
    skipped: int = 0
    exceptions: int = 0
    row_indexes: list[int] = field(default_factory=list)   # internal only; never serialized

    def summary(self) -> dict[str, Any]:
        return {"name": self.name, "checked": self.checked, "failed": self.failed,
                "skipped": self.skipped, "exceptions": self.exceptions}


@dataclass
class ImportResult:
    ok: bool
    status: Literal["staged", "failed"]
    failed_step: str | None
    source_key: str
    audience: Audience
    file: str
    sheet: str
    layout: str
    records: list[NormalizedRecord]
    report: ValidationReport
    verification: VerificationResult
    business_date: BusinessDate
    field_confidence: dict[str, Confidence]
    match_reports: list[MatchReport] = field(default_factory=list)
    rule_results: list[RuleResult] = field(default_factory=list)
    attribute_sources: tuple[AttributeSource, ...] = ()

    @property
    def confidence(self) -> Confidence | None:
        """Worst confidence over the fields present (None if no fields)."""
        states = list(self.field_confidence.values())
        return max(states, key=lambda c: c.rank) if states else None

    def summary(self) -> dict[str, Any]:
        """Safe to print, log or return from a future API: no records, no values, no paths."""
        return {
            "ok": self.ok,
            "status": self.status,
            "failed_step": self.failed_step,
            "source": self.source_key,
            "audience": self.audience.value,
            "file": self.file,
            "sheet": self.sheet,
            "layout": self.layout,
            "record_count": len(self.records),
            "business_date": self.business_date.model_dump(mode="json"),
            "field_confidence": {k: v.value for k, v in self.field_confidence.items()},
            "report": self.report.to_dict(),
            "verification": self.verification.model_dump(mode="json"),
            "match_reports": [m.model_dump(mode="json") for m in self.match_reports],
            "rules": [r.summary() for r in self.rule_results],
            "attribute_sources": [
                {"attribute": a.attribute, "master": a.master, "master_field": a.master_field,
                 "master_key": a.master_key, "via_role": a.via_role.value} for a in self.attribute_sources],
        }


@dataclass
class _Raw:
    row: int
    cells: dict[str, tuple[FieldSpec, CellRead]]      # record key → (normalization spec, cell)
    labels: dict[str, str | None] = field(default_factory=dict)
    region: dict[str, str] = field(default_factory=dict)  # accumulator name → column letter / block


# ---------------------------------------------------------------------------------------------
def run_import(path: str | Path, spec: SourceSpec, *, declared_date: dt.date | None = None,
               masters: Mapping[str, MasterIndex] | None = None) -> ImportResult:
    spec.check()
    path = Path(path)
    basename = path.name
    report = ValidationReport(source=spec.key, audience=spec.audience, file=basename,
                              sheet=spec.sheet_label(), layout=spec.layout_name)
    state = _State(spec=spec, report=report)

    if not spec.matches_filename(basename):
        report.add(Severity.ERROR, "filename_mismatch", "The file name does not match the source's declared patterns.")
    if path.suffix.lower() not in SUPPORTED_SUFFIXES:
        report.add(Severity.ERROR, "unsupported_format", "Only .xlsx / .xlsm workbooks are supported.")
        return state.fail("validate")
    try:
        links = external_link_targets(path)
    except (zipfile.BadZipFile, ParseError, KeyError, OSError):
        report.add(Severity.ERROR, "unreadable_workbook", "The workbook could not be opened as an .xlsx package.")
        return state.fail("validate")
    if links:
        report.external_links.count = len(links)
        report.external_links.names = sorted(set(links))
        report.add(spec.external_links, "external_link",
                   f"The workbook links to {len(links)} external workbook(s); linked files are not opened.",
                   count=len(links))

    try:
        wb_ctx = WorkbookPair(path).__enter__()
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError):
        report.add(Severity.ERROR, "unreadable_workbook", "The workbook could not be read.")
        return state.fail("validate")
    try:
        sheet = wb_ctx.sheet(spec.sheet)
        if sheet is None:
            report.add(Severity.ERROR, "sheet_missing", "The declared sheet does not exist.")
            return state.fail("validate")
        raws = (_read_table(sheet, spec, report) if isinstance(spec.layout, TableLayout)
                else _read_matrix(sheet, spec, report))
        if raws is None:
            return state.fail("validate")
        state.normalize(raws)
        state.run_rules()
        state.business_date = resolve_business_date(spec.business_date, report=report, records=state.records,
                                                    sheet=sheet, basename=basename, declared=declared_date)
        report.business_date = state.business_date
    finally:
        wb_ctx.__exit__(None, None, None)

    if report.errors:
        return state.fail("validate")
    if report.warnings and not spec.allow_warnings:
        report.add(Severity.ERROR, "warnings_not_allowed", "This source does not allow warnings.")
        return state.fail("validate")
    state.resolve_entities(masters or {})
    if report.errors:
        return state.fail("resolve")
    state.verify()
    if not state.verification.ok:
        failed = [c.name for c in state.verification.checks if not c.ok]
        report.add(Severity.ERROR, "verification_failed", "Verification failed: " + ", ".join(failed) + ".",
                   count=len(failed))
        return state.fail("verify")
    return state.result(ok=True)


# ---------------------------------------------------------------------------------------------
def _read_table(sheet: SheetPair, spec: SourceSpec, report: ValidationReport) -> list[_Raw] | None:
    layout: TableLayout = spec.layout  # type: ignore[assignment]
    hm = resolve_headers(sheet.row(layout.header_row), spec.fields)
    for sev, code, fname, region, msg in hm.problems:
        report.add(sev, code, msg, field=fname, region=region)
    for raw_name, cols in hm.duplicates.items():
        letters = ",".join(get_column_letter(c) for c in cols)
        users = [f.name for f in spec.fields if f.header.name == raw_name]  # type: ignore[union-attr]
        report.add(Severity.INFO, "duplicate_header",
                   f"A header name appears {len(cols)} times in the header row (columns {letters}).",
                   field=users[0] if users else None, region=letters, count=len(cols))
    if hm.unmapped:
        report.add(Severity.INFO, "unmapped_columns", f"{hm.unmapped} header column(s) are not used by the spec.",
                   count=hm.unmapped)
    if any(sev == Severity.ERROR for sev, *_ in hm.problems):
        return None
    if not hm.columns:
        report.add(Severity.ERROR, "no_fields", "No declared field was found at the header row.")
        return None
    fields = {f.name: f for f in spec.fields if f.name in hm.columns}
    letters = {name: get_column_letter(c) for name, c in hm.columns.items()}
    max_col = max(hm.columns.values())
    raws: list[_Raw] = []
    for r, cells in sheet.iter_rows(min_row=layout.first_data_row, max_col=max_col):
        sel = {name: (cells[c - 1] if c - 1 < len(cells) else EMPTY_CELL) for name, c in hm.columns.items()}
        if all(c.kind == CellKind.EMPTY for c in sel.values()):
            if layout.end == "first_blank":
                break
            report.blank_rows_skipped += 1
            continue
        raws.append(_Raw(row=r, cells={n: (fields[n], c) for n, c in sel.items()}, region=letters))
    report.rows_read = len(raws)
    return raws


def _block_spec(block: MatrixRange | MatrixSection) -> FieldSpec:
    return FieldSpec(name=block.name, type=block.type, required=block.required, confidence=block.confidence,
                     tolerate_invalid_cells=block.tolerate_invalid_cells)


def _coord(r: int, c: int) -> str:
    return f"{get_column_letter(c)}{r}"


def _read_matrix(sheet: SheetPair, spec: SourceSpec, report: ValidationReport) -> list[_Raw] | None:
    layout: MatrixLayout = spec.layout  # type: ignore[assignment]
    min_col, min_row, max_col, max_row = range_boundaries(layout.bounds)
    grid = sheet.grid(min_row, min_col, max_row, max_col)
    at: dict[str, tuple[int, int]] = {}
    for a in layout.anchors:
        s_min_col, s_min_row, s_max_col, s_max_row = range_boundaries(a.search)
        hits = sorted((r, c) for (r, c), cell in grid.items()
                      if s_min_row <= r <= s_max_row and s_min_col <= c <= s_max_col and label_text(cell) == a.label)
        expected = cell_rc(a.expected)
        if not hits:
            report.add(Severity.ERROR, "anchor_missing", "Expected label not found in its search range.",
                       field=a.name, region=a.expected)
        elif len(hits) > 1:
            report.add(Severity.ERROR, "anchor_ambiguous", f"Expected label found {len(hits)} times.",
                       field=a.name, region=",".join(_coord(*h) for h in hits), count=len(hits))
        elif hits[0] != expected:
            report.add(Severity.ERROR, "anchor_moved", "Expected label is not at its declared cell.",
                       field=a.name, region=f"{a.expected}->{_coord(*hits[0])}")
        else:
            at[a.name] = expected
    if report.errors:
        return None

    def in_bounds(r: int, c: int) -> bool:
        return min_row <= r <= max_row and min_col <= c <= max_col

    raws: list[_Raw] = []
    cells_rec = _Raw(row=0, cells={})
    for f in spec.fields:
        if f.cell:
            r, c = cell_rc(f.cell)
        else:
            ar, ac = at[f.anchor]  # type: ignore[index]
            r, c = ar + f.offset[0], ac + f.offset[1]
        if not in_bounds(r, c):
            report.add(Severity.ERROR, "cell_out_of_bounds", "Declared cell lies outside the matrix bounds.",
                       field=f.name, region=_coord(r, c))
            continue
        cells_rec.cells[f.name] = (f, grid.get((r, c), EMPTY_CELL))
        cells_rec.region[f.name] = _coord(r, c)
    if cells_rec.cells:
        raws.append(cells_rec)

    for rg in layout.ranges:
        bspec = _block_spec(rg)
        c0, r0, c1, r1 = range_boundaries(rg.range)
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                raws.append(_Raw(row=r, cells={"value": (bspec, grid.get((r, c), EMPTY_CELL))},
                                 labels={"block": rg.name, "row_label": None, "column_label": _coord(r, c)},
                                 region={rg.name: rg.range}))

    for sec in layout.sections:
        bspec = _block_spec(sec)
        ar, ac = at[sec.anchor]
        drift = 0
        for j, lab in enumerate(sec.column_labels):
            if label_text(grid.get((ar, ac + 1 + j), EMPTY_CELL)) != lab:
                drift += 1
        if grid.get((ar, ac + 1 + len(sec.column_labels))) is not None:
            drift += 1
        if drift:
            report.add(Severity.ERROR, "section_label_mismatch",
                       f"{drift} column label cell(s) differ from the declared section layout.",
                       field=sec.name, region=_coord(ar, ac), count=drift)
            continue
        limit = sec.rows if sec.rows is not None else sec.max_rows
        labels_seen: dict[str, int] = {}
        section_raws: list[_Raw] = []
        missing = 0
        for i in range(1, limit + 1):
            r = ar + i
            if not in_bounds(r, ac):
                missing += 1 if sec.rows is not None else 0
                break
            row_label = label_text(grid.get((r, ac), EMPTY_CELL))
            if row_label in (None, ""):
                if sec.rows is not None:
                    missing += 1
                    continue
                break
            labels_seen[row_label] = labels_seen.get(row_label, 0) + 1
            for j, lab in enumerate(sec.column_labels):
                section_raws.append(_Raw(row=r, cells={"value": (bspec, grid.get((r, ac + 1 + j), EMPTY_CELL))},
                                         labels={"block": sec.name, "row_label": row_label, "column_label": lab},
                                         region={sec.name: _coord(ar, ac)}))
        dup = sum(n - 1 for n in labels_seen.values() if n > 1)
        if missing:
            report.add(Severity.ERROR, "section_row_missing", f"{missing} declared section row(s) have no row label.",
                       field=sec.name, region=_coord(ar, ac), count=missing)
        if dup:
            report.add(Severity.ERROR, "section_duplicate_row_label", f"{dup} repeated row label(s) in the section.",
                       field=sec.name, region=_coord(ar, ac), count=dup)
        if not section_raws and sec.required:
            report.add(Severity.ERROR, "section_empty", "The section has no rows.", field=sec.name,
                       region=_coord(ar, ac))
        raws.extend(section_raws)
    if report.errors:
        return None
    report.rows_read = len(raws)
    return raws


# ---------------------------------------------------------------------------------------------
@dataclass
class _State:
    spec: SourceSpec
    report: ValidationReport
    records: list[NormalizedRecord] = field(default_factory=list)
    accs: dict[str, FieldAccumulator] = field(default_factory=dict)
    field_confidence: dict[str, Confidence] = field(default_factory=dict)
    declared: dict[str, Confidence] = field(default_factory=dict)   # record value key → declared state
    business_date: BusinessDate = field(default_factory=BusinessDate)
    verification: VerificationResult = field(default_factory=VerificationResult)
    match_reports: list[MatchReport] = field(default_factory=list)
    rule_results: list[RuleResult] = field(default_factory=list)
    raw_count: int = 0

    # -- normalize ----------------------------------------------------------------------------
    def normalize(self, raws: list[_Raw]) -> None:
        spec, report = self.spec, self.report
        self.raw_count = len(raws)
        regions: dict[str, str] = {}
        staged: list[tuple[int, dict[str, NormalizedValue]]] = []
        ext_refs = 0
        for raw in raws:
            values: dict[str, NormalizedValue] = {}
            for key, label in raw.labels.items():
                values[key] = NormalizedValue(label, spec.default_confidence, CellKind.LITERAL)
            for key, (fspec, cell) in raw.cells.items():
                declared = spec.confidence_of(fspec)
                acc = self.accs.get(fspec.name)
                if acc is None:
                    acc = self.accs[fspec.name] = FieldAccumulator(fspec, declared)
                    regions[fspec.name] = raw.region.get(fspec.name, raw.region.get(key, ""))
                n = normalize_cell(fspec, cell, declared)
                acc.add(cell, n)
                ext_refs += cell.external_ref
                values[key] = n.nv
                self.declared[key if key != "value" else f"value@{fspec.name}"] = declared
            staged.append((raw.row, values))

        if ext_refs:
            report.external_links.formula_refs = ext_refs
            report.add(spec.external_links, "external_formula_ref",
                       f"{ext_refs} formula cell(s) reference another workbook; values are Excel's cached results.",
                       count=ext_refs)

        # Field-level policy → issues and field confidence.
        unusable: set[str] = set()
        for name, acc in self.accs.items():
            f, conf, region = acc.spec, acc.declared, regions.get(name) or None
            strict = f.required and not f.tolerate_invalid_cells
            cell_sev = Severity.ERROR if strict else Severity.WARNING
            if acc.excel_error_count:
                detail = ", ".join(f"{k} ×{v}" for k, v in sorted(acc.errors.items()))
                report.add(cell_sev, "excel_error", f"{acc.excel_error_count} Excel error cell(s): {detail}.",
                           field=name, region=region, count=acc.excel_error_count)
            for code, text in ((TYPE_MISMATCH, "do not have the declared type"),
                               (ID_FLOAT, "are identifiers stored as non-integer or oversized numbers"),
                               (ID_MALFORMED, "do not match the declared identifier pattern")):
                n = acc.invalid.get(code, 0)
                if n:
                    report.add(cell_sev, code, f"{n} cell(s) {text}.", field=name, region=region, count=n)
            if acc.cell_invalid_count and not f.required and not f.tolerate_invalid_cells:
                conf = Confidence.INVALID
                unusable.add(name)
                report.add(Severity.WARNING, "field_unusable",
                           "Optional field has invalid cells; the whole field is marked invalid for this import.",
                           field=name, region=region, count=acc.cell_invalid_count)
            n = acc.invalid.get(ID_PLACEHOLDER, 0)
            if n:
                report.add(f.on_placeholder, ID_PLACEHOLDER,
                           f"{n} placeholder identifier(s); they are marked invalid and never joined.",
                           field=name, region=region, count=n)
            if acc.missing_cache:
                sev = f.missing_cache or spec.missing_cache
                report.add(sev, "formula_missing_cache",
                           f"{acc.missing_cache} formula cell(s) have no saved result; the value is unknown.",
                           field=name, region=region, count=acc.missing_cache)
            if acc.formula:
                if f.formulas == "forbid":
                    report.add(Severity.ERROR, "formula_forbidden", f"{acc.formula} formula cell(s) in a literal-only field.",
                               field=name, region=region, count=acc.formula)
                elif f.formulas == "warn":
                    report.add(Severity.WARNING, "formula_present", f"{acc.formula} formula cell(s).",
                               field=name, region=region, count=acc.formula)
                else:
                    report.add(Severity.INFO, "formula_cached_values",
                               f"{acc.formula} value(s) come from Excel's last saved calculation (not recalculated).",
                               field=name, region=region, count=acc.formula)
            if acc.notes.get("id_numeric_cell"):
                n = acc.notes["id_numeric_cell"]
                report.add(Severity.WARNING, "id_numeric_cell",
                           f"{n} identifier(s) are stored as numbers; leading zeros cannot be verified.",
                           field=name, region=region, count=n)
            nulls = acc.rows - acc.filled
            if not f.nullable and nulls:
                report.add(Severity.ERROR if f.required else Severity.WARNING, "null_values",
                           f"{nulls} empty cell(s) in a non-nullable field.", field=name, region=region, count=nulls)
            self.field_confidence[name] = conf

        # Optional table fields whose header is missing are present as `invalid`.
        for issue in report.issues:
            if issue.code == "header_missing_optional" and issue.field:
                self.field_confidence[issue.field] = Confidence.INVALID

        for row, values in staged:
            if unusable:
                for key, nv in list(values.items()):
                    owner = key if key in unusable else None
                    if key == "value" and values.get("block") is not None and values["block"].value in unusable:
                        owner = values["block"].value
                    if owner is not None:
                        values[key] = NormalizedValue(None, Confidence.INVALID, nv.kind, nv.error or "field_unusable")
            self.records.append(NormalizedRecord(key="", row=row, values=values))

        self._assign_keys()
        self.report.fields = [acc.stats(self.field_confidence[n]) for n, acc in self.accs.items()]

    def _assign_keys(self) -> None:
        spec = self.spec
        if isinstance(spec.layout, MatrixLayout):
            identities = [("cells",) if r.row == 0 and "block" not in r.values else
                          (r.get("block"), r.get("row_label"), r.get("column_label")) for r in self.records]
        else:
            identities = [tuple(r.get(f) for f in spec.identity_fields) for r in self.records]
            incomplete = sum(1 for ident in identities if spec.identity_fields and any(p is None for p in ident))
            if incomplete:
                self.report.add(Severity.WARNING, "identity_incomplete",
                                f"{incomplete} record(s) have an empty or invalid identity part.", count=incomplete)
        for rec, key in zip(self.records, keys_with_occurrence(spec.key, identities)):
            rec.key = key

    # -- rules --------------------------------------------------------------------------------
    def run_rules(self) -> None:
        for rule in self.spec.rules:
            res = RuleResult(rule.name)
            for i, rec in enumerate(self.records):
                vals = {f: rec.values.get(f) for f in rule.fields}
                if any(nv is None or nv.value is None or nv.confidence == Confidence.INVALID for nv in vals.values()):
                    res.skipped += 1
                    continue
                res.checked += 1
                try:
                    passed = bool(rule.check({f: nv.value for f, nv in vals.items()}))  # type: ignore[union-attr]
                except Exception:  # a rule must never crash the import or leak values in a traceback
                    res.exceptions += 1
                    passed = False
                if not passed:
                    res.failed += 1
                    res.row_indexes.append(i)
            self.rule_results.append(res)
            label = f"rule:{rule.name}"
            if res.failed:
                self.report.add(rule.severity, "rule_failed",
                                f"Rule failed on {res.failed} of {res.checked} checked row(s).",
                                field=label, count=res.failed)
            if res.skipped:
                self.report.add(Severity.INFO, "rule_skipped",
                                f"{res.skipped} row(s) not checked (empty or invalid inputs).",
                                field=label, count=res.skipped)

    # -- entity resolution ----------------------------------------------------------------------
    def resolve_entities(self, masters: Mapping[str, MasterIndex]) -> None:
        for f in self.spec.fields:
            if not f.master or f.name not in self.field_confidence:
                continue
            master = masters.get(f.master)
            if master is None:
                self.report.add(Severity.INFO, "er_master_not_supplied",
                                "The master index for this identifier was not supplied; not resolved.", field=f.name)
                continue
            resolutions, mr = resolve_values([r.values.get(f.name) for r in self.records], master,
                                             field_name=f.name, role=f.role)  # type: ignore[arg-type]
            for rec, res in zip(self.records, resolutions):
                rec.links[f.name] = res
            self.match_reports.append(mr)
            if mr.master_duplicate_keys:
                self.report.add(Severity.WARNING, "master_duplicate_keys",
                                f"The master has {mr.master_duplicate_keys} non-unique key(s); rows joining them "
                                "are reported as duplicate_master, never matched.",
                                field=f.name, count=mr.master_duplicate_keys)

    # -- verify -------------------------------------------------------------------------------
    def verify(self) -> None:
        v, spec = self.verification, self.spec
        v.add("row_count", len(self.records) == self.raw_count, f"{len(self.records)} of {self.raw_count}")
        keys = [r.key for r in self.records]
        v.add("keys_unique", len(keys) == len(set(keys)) and all(keys), f"{len(set(keys))} distinct")
        bad_err = sum(1 for r in self.records for nv in r.values.values()
                      if nv.error is not None and (nv.confidence != Confidence.INVALID or nv.value is not None))
        v.add("errors_marked_invalid", bad_err == 0, f"{bad_err} violation(s)")
        upgraded = 0
        for r in self.records:
            for key, nv in r.values.items():
                if key in ("block", "row_label", "column_label"):
                    continue
                owner = key if key != "value" else f"value@{r.values['block'].value}"
                declared = self.declared.get(owner)
                if declared is not None and nv.confidence.rank < declared.rank:
                    upgraded += 1
        v.add("confidence_never_upgraded", upgraded == 0, f"{upgraded} violation(s)")
        v.add("min_rows", len(self.records) >= spec.verify.min_rows, f"{len(self.records)} record(s)")
        for fname, limit in spec.verify.max_null_rate.items():
            acc = self.accs.get(fname)
            rate = (acc.rows - acc.filled) / acc.rows if acc and acc.rows else 1.0
            v.add(f"max_null_rate:{fname}", rate <= limit, f"{rate:.4f} <= {limit}")
        if spec.business_date.strategy != "none" and spec.verify.require_business_date:
            v.add("business_date", self.business_date.value is not None)

    # -- results ------------------------------------------------------------------------------
    def result(self, ok: bool, failed_step: str | None = None) -> ImportResult:
        return ImportResult(
            ok=ok, status="staged" if ok else "failed", failed_step=failed_step,
            source_key=self.spec.key, audience=self.spec.audience, file=self.report.file,
            sheet=self.spec.sheet_label(), layout=self.spec.layout_name,
            records=self.records if ok else [], report=self.report, verification=self.verification,
            business_date=self.business_date, field_confidence=dict(self.field_confidence),
            match_reports=self.match_reports if ok else [], rule_results=self.rule_results,
            attribute_sources=self.spec.attribute_sources)

    def fail(self, step: str) -> ImportResult:
        return self.result(ok=False, failed_step=step)
