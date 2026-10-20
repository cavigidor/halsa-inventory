"""Declarative source specifications.

A `SourceSpec` says exactly what a workbook must look like. The framework never guesses:
headers are matched byte-for-byte at a declared row, duplicate headers are selected by
(raw name, occurrence), matrix values are read from declared cells or label anchors, and any
drift is a validation issue.
"""
from __future__ import annotations

import fnmatch
import re
import unicodedata
from dataclasses import dataclass, field
from enum import Enum
from typing import Any, Callable, Literal

from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, range_boundaries

from .models import AttributeSource, Audience, Confidence, IdentifierRole, Severity


class SpecError(ValueError):
    """The spec itself is wrong (a programming error, not a data problem)."""


class FieldType(str, Enum):
    STRING = "string"
    OPAQUE_ID = "opaque_id"   # codes: always strings, never numbers
    INTEGER = "integer"
    DECIMAL = "decimal"
    DATE = "date"
    BOOLEAN = "boolean"
    ANY = "any"               # raw cell value, no conversion (validated-only regions with mixed content)


@dataclass(frozen=True)
class Header:
    """Exact raw header selector. `occurrence` (1-based) is required when the name is duplicated
    in the header row; with `occurrence=None` a duplicated name is an `ambiguous_header` error."""

    name: str
    occurrence: int | None = None


@dataclass(frozen=True)
class Anchor:
    """A label that must appear exactly once in `search`, at `expected`.

    Moved (found once elsewhere), missing, or ambiguous (found more than once) are errors.
    """

    name: str
    label: str
    expected: str     # e.g. "A5"
    search: str       # e.g. "A1:Z60"; must lie within the matrix bounds


@dataclass(frozen=True)
class FieldSpec:
    name: str
    type: FieldType = FieldType.STRING
    required: bool = True
    nullable: bool = True
    confidence: Confidence | None = None      # None → the spec's default_confidence
    # --- where (exactly one, depending on layout) ---
    header: Header | str | None = None        # table layout
    cell: str | None = None                   # matrix: explicit cell
    anchor: str | None = None                 # matrix: anchor name …
    offset: tuple[int, int] = (0, 1)          # … plus (rows, cols) offset from the anchor
    # --- identifiers ---
    trim: bool = False                        # strip whitespace only when declared
    pattern: str | None = None                # opaque id: valid codes fullmatch this regex
    placeholders: tuple[str, ...] = ()        # regexes for placeholder codes (reported, never joined)
    on_placeholder: Severity = Severity.WARNING
    role: IdentifierRole | None = None
    master: str | None = None                 # name of the master index to resolve against
    # --- policies ---
    missing_cache: Severity | None = None     # None → spec.missing_cache
    formulas: Literal["allow", "warn", "forbid"] = "allow"
    tolerate_invalid_cells: bool = False      # True: invalid cells are cell-level warnings, field stays usable

    def __post_init__(self) -> None:
        if isinstance(self.header, str):
            object.__setattr__(self, "header", Header(self.header))


@dataclass(frozen=True)
class TableLayout:
    header_row: int
    data_start_row: int | None = None   # default header_row + 1
    end: Literal["sheet_end", "first_blank"] = "sheet_end"
    data_end_row: int | None = None     # last data row (inclusive) when a layout has trailing control rows

    @property
    def first_data_row(self) -> int:
        return self.data_start_row or self.header_row + 1


@dataclass(frozen=True)
class MatrixRange:
    """Explicit rectangular range; one long record per cell (block, column_label=cell, value)."""

    name: str
    range: str
    type: FieldType = FieldType.DECIMAL
    confidence: Confidence | None = None
    required: bool = True
    tolerate_invalid_cells: bool = False


@dataclass(frozen=True)
class MatrixSection:
    """A labeled block anchored at `anchor`.

    The anchor's row holds the column labels (must equal `column_labels` exactly, starting one
    column right of the anchor). Row labels are in the anchor's column below it. Rows are read
    until the first empty row label, or for exactly `rows` rows when set. Long records:
    (block, row_label, column_label, value).
    """

    name: str
    anchor: str
    column_labels: tuple[str, ...]
    type: FieldType = FieldType.DECIMAL
    rows: int | None = None
    max_rows: int = 500
    confidence: Confidence | None = None
    required: bool = True
    tolerate_invalid_cells: bool = False


@dataclass(frozen=True)
class MatrixLayout:
    bounds: str                                   # explicit read window, e.g. "A1:N80"
    anchors: tuple[Anchor, ...] = ()
    ranges: tuple[MatrixRange, ...] = ()
    sections: tuple[MatrixSection, ...] = ()


@dataclass(frozen=True)
class DateRule:
    """Business-date strategy. Never file modification time."""

    strategy: Literal["none", "column_max", "snapshot_cell", "bucket_headers", "declared", "filename"] = "none"
    field: str | None = None
    cell: str | None = None
    cells: str | None = None
    filename_regex: str | None = None   # named groups: year, month, optional day

    @classmethod
    def column_max(cls, field: str) -> "DateRule":
        return cls("column_max", field=field)

    @classmethod
    def snapshot_cell(cls, cell: str) -> "DateRule":
        return cls("snapshot_cell", cell=cell)

    @classmethod
    def bucket_headers(cls, cells: str) -> "DateRule":
        return cls("bucket_headers", cells=cells)

    @classmethod
    def declared(cls) -> "DateRule":
        return cls("declared")

    @classmethod
    def from_filename(cls, regex: str) -> "DateRule":
        return cls("filename", filename_regex=regex)


@dataclass(frozen=True)
class RowRule:
    """Deterministic row-level rule. `check` gets {field: value} for rows where every listed field
    is valid and non-null; it returns True when the row passes."""

    name: str
    fields: tuple[str, ...]
    check: Callable[[dict[str, Any]], bool]
    severity: Severity = Severity.ERROR
    description: str = ""

    @classmethod
    def sum_identity(cls, name: str, total: str, parts: tuple[str, ...], *, tolerance: float = 0.005,
                     severity: Severity = Severity.ERROR) -> "RowRule":
        """total == sum(parts) within an absolute tolerance."""

        def check(v: dict[str, Any]) -> bool:
            return abs(float(v[total]) - sum(float(v[p]) for p in parts)) <= tolerance

        return cls(name, (total, *parts), check, severity, f"{total} = " + " + ".join(parts))


@dataclass(frozen=True)
class VerifySpec:
    min_rows: int = 0
    max_null_rate: dict[str, float] = field(default_factory=dict)
    require_business_date: bool = True   # only applies when the strategy is not "none"


@dataclass(frozen=True)
class SourceSpec:
    key: str
    audience: Audience
    filename_patterns: tuple[str, ...]
    sheet: str | int
    layout: TableLayout | MatrixLayout
    fields: tuple[FieldSpec, ...]
    default_confidence: Confidence = Confidence.UNVERIFIED
    business_date: DateRule = field(default_factory=DateRule)
    identity_fields: tuple[str, ...] = ()
    rules: tuple[RowRule, ...] = ()
    verify: VerifySpec = field(default_factory=VerifySpec)
    allow_warnings: bool = True
    missing_cache: Severity = Severity.ERROR
    external_links: Severity = Severity.WARNING
    attribute_sources: tuple[AttributeSource, ...] = ()
    description: str = ""

    # -- helpers ------------------------------------------------------------------------------
    @property
    def layout_name(self) -> str:
        return "table" if isinstance(self.layout, TableLayout) else "matrix"

    def field(self, name: str) -> FieldSpec:
        for f in self.fields:
            if f.name == name:
                return f
        raise KeyError(name)

    def confidence_of(self, f: FieldSpec | MatrixRange | MatrixSection) -> Confidence:
        return f.confidence or self.default_confidence

    def sheet_label(self) -> str:
        return self.sheet if isinstance(self.sheet, str) else f"#{self.sheet}"

    def matches_filename(self, basename: str) -> bool:
        """Case-sensitive glob match after Unicode NFC normalization (macOS stores NFD names)."""
        name = unicodedata.normalize("NFC", basename)
        return any(fnmatch.fnmatchcase(name, unicodedata.normalize("NFC", p)) for p in self.filename_patterns)

    def check(self) -> None:
        """Raise SpecError if the spec is internally inconsistent."""
        problems: list[str] = []
        if not self.key or not re.fullmatch(r"[a-z0-9_]+", self.key):
            problems.append("key must be lowercase [a-z0-9_]")
        if not self.filename_patterns:
            problems.append("at least one filename pattern is required")
        names = [f.name for f in self.fields]
        if len(names) != len(set(names)):
            problems.append("duplicate field names")
        reserved = {"block", "row_label", "column_label", "value"}
        table = isinstance(self.layout, TableLayout)
        if table:
            if self.layout.header_row < 1 or self.layout.first_data_row <= self.layout.header_row:
                problems.append("invalid header/data rows")
            if self.layout.data_end_row is not None and self.layout.data_end_row < self.layout.first_data_row:
                problems.append("data_end_row before the first data row")
            for f in self.fields:
                if not isinstance(f.header, Header) or f.cell or f.anchor:
                    problems.append(f"table field {f.name} needs exactly a header")
                elif f.header.occurrence is not None and f.header.occurrence < 1:
                    problems.append(f"field {f.name}: occurrence is 1-based")
        else:
            try:
                range_boundaries(self.layout.bounds)
            except (ValueError, TypeError):
                problems.append("matrix bounds must be a range like A1:N80")
            anchor_names = {a.name for a in self.layout.anchors}
            if len(anchor_names) != len(self.layout.anchors):
                problems.append("duplicate anchor names")
            for f in self.fields:
                if f.header is not None or (f.cell is None) == (f.anchor is None):
                    problems.append(f"matrix field {f.name} needs exactly one of cell / anchor")
                if f.anchor and f.anchor not in anchor_names:
                    problems.append(f"field {f.name}: unknown anchor")
                if f.name in reserved:
                    problems.append(f"field name {f.name} is reserved")
            blocks = [b.name for b in (*self.layout.ranges, *self.layout.sections)]
            if len(blocks) != len(set(blocks)) or set(blocks) & set(names):
                problems.append("range/section names must be unique and differ from field names")
            try:
                b = range_boundaries(self.layout.bounds)
                for a in self.layout.anchors:
                    s_ = range_boundaries(a.search)
                    r, c = cell_rc(a.expected)
                    if not (b[0] <= s_[0] and b[1] <= s_[1] and s_[2] <= b[2] and s_[3] <= b[3]):
                        problems.append(f"anchor {a.name}: search range outside bounds")
                    if not (s_[0] <= c <= s_[2] and s_[1] <= r <= s_[3]):
                        problems.append(f"anchor {a.name}: expected cell outside its search range")
                for rg in self.layout.ranges:
                    s_ = range_boundaries(rg.range)
                    if not (b[0] <= s_[0] and b[1] <= s_[1] and s_[2] <= b[2] and s_[3] <= b[3]):
                        problems.append(f"range {rg.name}: outside bounds")
            except (ValueError, TypeError):
                problems.append("invalid anchor/range reference")
            for s in self.layout.sections:
                if s.anchor not in anchor_names:
                    problems.append(f"section {s.name}: unknown anchor")
                if not s.column_labels:
                    problems.append(f"section {s.name}: column labels required")
        for f in self.fields:
            for rx in (f.pattern, *f.placeholders):
                if rx is not None:
                    try:
                        re.compile(rx)
                    except re.error:
                        problems.append(f"field {f.name}: bad regex")
            if (f.pattern or f.placeholders) and f.type != FieldType.OPAQUE_ID:
                problems.append(f"field {f.name}: pattern/placeholders apply to opaque_id fields only")
            if f.master and not f.role:
                problems.append(f"field {f.name}: an identifier resolved against a master needs a role")
        for idf in self.identity_fields:
            if idf not in names:
                problems.append(f"identity field {idf} is not a field")
        for r in self.rules:
            for rf in r.fields:
                if rf not in names:
                    problems.append(f"rule {r.name}: unknown field {rf}")
        bd = self.business_date
        if bd.strategy == "column_max" and (bd.field not in names or self.field(bd.field).type != FieldType.DATE):
            problems.append("column_max needs a date field")
        if bd.strategy == "snapshot_cell" and not _is_cell(bd.cell):
            problems.append("snapshot_cell needs a cell")
        if bd.strategy == "bucket_headers":
            try:
                range_boundaries(bd.cells or "")
            except (ValueError, TypeError):
                problems.append("bucket_headers needs a range")
        if bd.strategy == "filename":
            if not bd.filename_regex or not {"year", "month"} <= set(re.compile(bd.filename_regex).groupindex):
                problems.append("filename strategy needs a regex with year and month groups")
        if problems:
            raise SpecError(f"{self.key}: " + "; ".join(problems))


def _is_cell(ref: str | None) -> bool:
    if not ref:
        return False
    try:
        coordinate_from_string(ref)
        return True
    except ValueError:
        return False


def cell_rc(ref: str) -> tuple[int, int]:
    """'B7' → (7, 2)."""
    col, row = coordinate_from_string(ref)
    return row, column_index_from_string(col)
