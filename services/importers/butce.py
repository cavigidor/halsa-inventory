"""M6-6 — Bütçe importer (owner / Yönetim). Docs: docs/m6_butce.md; decision D-041.

Source module on the M6-1 framework: one explicit matrix spec (`butce_spec`) read by `run_import`, orchestrated
by `import_butce`. No reader of its own; nothing recalculated, repaired or filled; owner audience only.

Verified layout (`BUTCE_<YYYY> <label>.xlsx`, sheet `<YYYY>`, bounds A1:ED141, the 2026 template):
- Seven stacked sales blocks (title rows 4, 23, 42, 61, 80, 99, 118). Each block: title row, segment row
  (`2018 Forecast`, `2018  GERCEK` … `<Y>  GERCEK`), unit row (`Months`, then per segment an amount column labelled
  `YTL`/`TL` and two exchange-rate columns `$ kur` / `€ kur`, whose order differs between Forecast and GERCEK),
  12 month rows (January … December), a `Total` row and three derived rows (growth ratios, notes).
- An expense area right of block 1 (BF4:ED19): 19 groups (`2016 expenses`, then `<y> expens forecast` /
  `<y> expenses`), each Finance / Other / Total, by month, plus a Total row. Above it, row 2 holds unlabelled
  annual inputs and row 1 holds cells linked to another workbook.
- Five `BÜTÇE ÖZET` panels (senaryo 1 / senaryo 2 × five lines) in CK24:ED31.
- Everything else (derived rows, deviation/returns panels, rate growth, notes) is validated and counted, not
  normalized. `Sayfa1` is a scratch summary and is not read.

Cell-level policy: a structural/header problem fails the import; an Excel error, a missing formula cache, a
non-numeric value or an external-workbook dependency makes only that cell invalid (excluded from usable coverage).
Valid values are `unverified`. File authority is unresolved: nothing here selects a file of record.
"""
from __future__ import annotations

import datetime as dt
import json
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from enum import Enum
from pathlib import Path
from typing import Any

from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries
from pydantic import BaseModel, ConfigDict, Field

from .base import ImportResult, run_import
from .excel import WorkbookPair
from .keys import record_key
from .models import Audience, CellKind, Confidence, NormalizedValue, Severity, ValidationReport, VerificationResult
from .spec import Anchor, FieldType, MatrixLayout, MatrixRange, SourceSpec

BUTCE_FILENAME = re.compile(r"^BUTCE_(?P<year>\d{4})[ _]+(?P<label>[^/\\]+?)\s*\.xlsx$")
SUPPORTED_YEARS = frozenset({2026})
FIRST_SEGMENT_YEAR = 2018
BOUNDS = "A1:ED141"
TOL = 0.01                       # absolute tolerance for validation arithmetic only
RATE_TOL = 1e-9
MONTHS_EN = ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
             "November", "December")
UNIT_LABELS = {"Y": "YTL", "T": "TL", "$": "$ kur", "E": "€ kur"}
SCRATCH_SHEETS = ("Sayfa1",)
EXTERNAL_DEPENDENCY = "external_dependency"
TYPE_MISMATCH = "type_mismatch"


def _ci(col: str) -> int:
    return column_index_from_string(col)


def _coord(col: int, row: int) -> str:
    return f"{get_column_letter(col)}{row}"


# ---------------------------------------------------------------------------------------------
# Template (2026)
# ---------------------------------------------------------------------------------------------
class BudgetBlock(str, Enum):
    """Structural identity of the seven sales blocks (fixed order). Not an accounting classification."""

    TOTAL_SALES = "total_sales"
    TOTAL_SALES_WITHOUT_TROTEC = "total_sales_without_trotec"
    GROUP_LIST_1 = "group_list_1"    # title lists several product/brand names (not committed)
    GROUP_LIST_2 = "group_list_2"    # title lists several product/brand names (not committed)
    TRODAT = "trodat"
    IMALAT = "imalat"
    TROTEC = "trotec"


@dataclass(frozen=True)
class BlockDef:
    index: int
    title_row: int
    block: BudgetBlock
    title: str | None    # exact AL/AR/AX title; None → a group-list title verified for consistency only
    units: str           # unit codes for one Forecast + GERCEK pair (repeats for every year)


BLOCKS = (
    BlockDef(1, 4, BudgetBlock.TOTAL_SALES, "TOTAL SALES", "Y$EYE$"),
    BlockDef(2, 23, BudgetBlock.TOTAL_SALES_WITHOUT_TROTEC, "TOTAL SALES WITHOUT TROTEC", "Y$EYE$"),
    BlockDef(3, 42, BudgetBlock.GROUP_LIST_1, None, "T$EYE$"),
    BlockDef(4, 61, BudgetBlock.GROUP_LIST_2, None, "T$ETE$"),
    BlockDef(5, 80, BudgetBlock.TRODAT, "TRODAT", "Y$ETE$"),
    BlockDef(6, 99, BudgetBlock.IMALAT, "İMALAT", "T$ETE$"),
    BlockDef(7, 118, BudgetBlock.TROTEC, "TROTEC", "T$ETE$"),
)
#: Title cells of a block row (merged ranges B:Y, Z:AK, AL:AQ, AR:AW, AX:BC) and the segment years they cover.
#: B and Z carry the historical title in every block; AL/AR/AX carry the block's own title.
HISTORICAL_TITLE = "TOTAL SALES"
TITLE_CELLS = (("B", FIRST_SEGMENT_YEAR, 2021), ("Z", 2022, 2023), ("AL", 2024, 2024), ("AR", 2025, 2025),
               ("AX", 2026, 2026))
BLOCK_TITLE_COLS = ("AL", "AR", "AX")


class Scenario(str, Enum):
    FORECAST = "forecast"          # 'Forecast' / 'expens forecast'
    ACTUAL = "actual"              # 'GERCEK' (gerçek = actual)
    UNQUALIFIED = "unqualified"    # 'YYYY expenses': the label carries no scenario word (not asserted to be actual)
    SENARYO_1 = "senaryo_1"        # BÜTÇE ÖZET panel scenario, source-native
    SENARYO_2 = "senaryo_2"


@dataclass(frozen=True)
class Segment:
    index: int
    col: int
    year: int
    scenario: Scenario
    label: str


def segments(year: int) -> list[Segment]:
    out = []
    for i in range(2 * (year - FIRST_SEGMENT_YEAR + 1)):
        y = FIRST_SEGMENT_YEAR + i // 2
        fc = i % 2 == 0
        out.append(Segment(i, 2 + 3 * i, y, Scenario.FORECAST if fc else Scenario.ACTUAL,
                           f"{y} Forecast" if fc else f"{y}  GERCEK"))   # two spaces before GERCEK, as in the source
    return out


def segment_units(b: BlockDef, s: Segment) -> tuple[str, str, str]:
    codes = b.units[(s.index % 2) * 3:(s.index % 2) * 3 + 3]
    return tuple(UNIT_LABELS[c] for c in codes)  # type: ignore[return-value]


def segment_title_cell(s: Segment) -> str:
    for col, y0, y1 in TITLE_CELLS:
        if y0 <= s.year <= y1:
            return col
    raise ValueError(s.year)


EXPENSE_FIRST_COL = _ci("BF")
EXPENSE_COMPONENTS = ("Finance", "Other", "Total")
EXPENSE_HEADER_ROW, EXPENSE_COMPONENT_ROW, EXPENSE_UNIT_ROW = 4, 5, 6
EXPENSE_FIRST_MONTH_ROW, EXPENSE_TOTAL_ROW = 7, 19
EXPENSE_INPUT_ROW, EXPENSE_LINKED_ROW = 2, 1


def expense_groups(year: int) -> list[tuple[int, Scenario, str]]:
    groups = [(2016, Scenario.UNQUALIFIED, "2016 expenses")]
    for y in range(FIRST_SEGMENT_YEAR, year + 1):
        groups += [(y, Scenario.FORECAST, f"{y} expens forecast"), (y, Scenario.UNQUALIFIED, f"{y} expenses")]
    return groups


PANEL_TITLE, PANEL_TITLE_ROW = "BÜTÇE ÖZET", 25
PANEL_TARGET_LABEL, PANEL_TARGET_ROW = "TOPLAM SATIŞ HEDEFİ", 24
PANEL_SCENARIO_ROW, PANEL_FIRST_LINE_ROW = 26, 27
PANEL_SCENARIOS = (("senaryo 1", Scenario.SENARYO_1), ("senaryo 2", Scenario.SENARYO_2))
PANEL_LINES = ("satış", "satış maliyeti", "kar", "Giderler")
#: (label column, panel year — equals the expense-group header above it, last line label, has target label)
PANELS = (("CL", 2021, "Kar", False), ("DB", 2023, "Kar", False), ("DJ", 2024, "EBIT", True),
          ("DR", 2025, "EBIT", True), ("DZ", 2026, "EBIT", True))


# ---------------------------------------------------------------------------------------------
# Spec
# ---------------------------------------------------------------------------------------------
def _labels(name: str, ref: str) -> MatrixRange:
    return MatrixRange(name, ref, FieldType.STRING, confidence=Confidence.TRUSTED, required=True,
                       tolerate_invalid_cells=False)


def _values(name: str, ref: str) -> MatrixRange:
    return MatrixRange(name, ref, FieldType.DECIMAL, confidence=Confidence.UNVERIFIED, required=True,
                       tolerate_invalid_cells=True)


def _validated(name: str, ref: str) -> MatrixRange:
    return MatrixRange(name, ref, FieldType.ANY, confidence=Confidence.UNVERIFIED, required=False,
                       tolerate_invalid_cells=True)


def butce_ranges(year: int) -> tuple[MatrixRange, ...]:
    """Non-overlapping ranges that together cover the whole bounds exactly."""
    r: list[MatrixRange] = [
        _validated("top_header", "A1:BC2"), _validated("top_spacer", "BD1:BE2"), _values("expense_inputs", "BF1:ED2"),
        _validated("spacer", "A3:ED3"),
    ]
    for b in BLOCKS:
        t, k = b.title_row, b.index
        r += [_labels(f"b{k}_title", f"A{t}:BC{t}"), _labels(f"b{k}_segments", f"A{t + 1}:BC{t + 1}"),
              _labels(f"b{k}_units", f"A{t + 2}:BC{t + 2}"), _labels(f"b{k}_row_labels", f"A{t + 3}:A{t + 15}"),
              _values(f"b{k}_months", f"B{t + 3}:BC{t + 14}"), _values(f"b{k}_total", f"B{t + 15}:BC{t + 15}"),
              _validated(f"b{k}_derived", f"A{t + 16}:BC{t + 18}")]
        if k == 1:
            r += [_validated("rate_growth", "BD4:BE22"), _labels("expense_headers", "BF4:ED6"),
                  _values("expense_months", "BF7:ED18"), _values("expense_total", "BF19:ED19"),
                  _validated("expense_derived", "BF20:ED22")]
        elif k == 2:
            r += [_validated("b2_right", f"BD{t}:CJ{t + 18}"), _validated("panel_top", "CK23:ED23"),
                  _validated("panels", "CK24:ED31"), _validated("panel_below", f"CK32:ED{t + 18}")]
        else:
            r.append(_validated(f"b{k}_right", f"BD{t}:ED{t + 18}"))
    r.append(_validated("notes", "A137:ED141"))
    return tuple(r)


def butce_spec(year: int) -> SourceSpec:
    anchors: list[Anchor] = []
    for b in BLOCKS:
        t, k = b.title_row, b.index
        anchors += [
            Anchor(f"b{k}_title", HISTORICAL_TITLE, expected=f"B{t}", search=f"B{max(1, t - 2)}:B{t + 2}"),
            Anchor(f"b{k}_months_header", "Months", expected=f"A{t + 2}", search=f"A{t}:A{t + 4}"),
            Anchor(f"b{k}_january", MONTHS_EN[0], expected=f"A{t + 3}", search=f"A{t + 1}:A{t + 5}"),
            Anchor(f"b{k}_total", "Total", expected=f"A{t + 15}", search=f"A{t + 13}:A{t + 17}"),
        ]
    groups = expense_groups(year)
    last_col = get_column_letter(EXPENSE_FIRST_COL + 4 * (len(groups) - 1))
    anchors += [Anchor("expense_first_group", groups[0][2], expected="BF4", search="BF2:BF8"),
                Anchor("expense_last_group", groups[-1][2], expected=f"{last_col}4", search=f"{last_col}2:{last_col}8")]
    anchors += [Anchor(f"panel_{col}", PANEL_TITLE, expected=f"{col}{PANEL_TITLE_ROW}", search=f"{col}23:{col}27")
                for col, *_ in PANELS]
    return SourceSpec(
        key="butce", audience=Audience.OWNER, filename_patterns=("BUTCE_*.xlsx",), sheet=str(year),
        layout=MatrixLayout(bounds=BOUNDS, anchors=tuple(anchors), ranges=butce_ranges(year)), fields=(),
        default_confidence=Confidence.UNVERIFIED, missing_cache=Severity.WARNING, external_links=Severity.WARNING,
        description="Bütçe planning workbook (owner only; cell-level validity).")


# ---------------------------------------------------------------------------------------------
# Models (in memory)
# ---------------------------------------------------------------------------------------------
class BudgetRegion(str, Enum):
    """Normalized fact regions."""

    SALES_MONTH = "sales_month"                    # block × segment × month (amount column)
    SALES_TOTAL = "sales_total"                    # block × segment, Total row (amount column)
    RATE_MONTH = "rate_month"                      # block 1 × segment × rate column × month
    EXPENSE_MONTH = "expense_month"                # expense group × component × month
    EXPENSE_TOTAL = "expense_total"                # expense group × component, Total row
    EXPENSE_ANNUAL_INPUT = "expense_annual_input"  # row 2 above an expense group (unlabelled)
    EXPENSE_LINKED_INPUT = "expense_linked_input"  # row 1 above an expense group (linked-workbook row)
    BUDGET_SUMMARY = "budget_summary"              # BÜTÇE ÖZET panel × scenario × line


class ValueType(str, Enum):
    AMOUNT = "amount"                  # reported in the source's local-currency label (YTL / TL), never converted
    CURRENCY_RATE = "currency_rate"    # `$ kur` / `€ kur`: an exchange-rate column (label kept verbatim)


class CellValidity(str, Enum):
    VALID = "valid"
    EXCEL_ERROR = "excel_error"
    MISSING_CACHE = "missing_cache"
    EXTERNAL_DEPENDENCY = "external_dependency"
    TYPE_MISMATCH = "type_mismatch"


class SourceStatus(str, Enum):
    COMPLETE = "complete"                  # no invalid cell anywhere (still not authoritative)
    PARTIALLY_USABLE = "partially_usable"  # usable facts plus invalid cells / external dependencies
    UNUSABLE = "unusable"                  # no usable fact


#: Validated-only regions (counted, never normalized), by range name.
_VALIDATED_REGION = {
    "top_header": "top_header", "top_spacer": "spacer", "spacer": "spacer", "rate_growth": "rate_growth",
    "expense_derived": "expense_derived_rows", "b2_right": "deviation_and_side_panels",
    "panel_top": "panel_side_cells", "panels": "panel_side_cells", "panel_below": "panel_side_cells",
    "notes": "notes", "expense_inputs": "expense_separators", "expense_months": "expense_separators",
    "expense_total": "expense_separators", "expense_headers": "labels",
}


def _validated_region(range_name: str) -> str:
    if range_name in _VALIDATED_REGION:
        return _VALIDATED_REGION[range_name]
    suffix = range_name.split("_", 1)[1]
    return {"title": "labels", "segments": "labels", "units": "labels", "row_labels": "labels",
            "months": "group_rate_columns", "total": "sales_total_rate_columns", "derived": "sales_derived_rows",
            "right": "deviation_and_side_panels"}[suffix]


@dataclass(frozen=True, slots=True)
class _Pos:
    region: BudgetRegion
    block: BudgetBlock | None
    year: int
    scenario: Scenario
    month: int | None
    value_type: ValueType
    source_label: str          # structural column label: YTL / TL / $ kur / € kur / component / panel title
    component: str | None = None
    line_label: str | None = None
    title_conflict: bool = False


@dataclass(frozen=True, slots=True)
class BudgetFact:
    record_key: str
    region: BudgetRegion
    block: BudgetBlock | None
    year: int
    scenario: Scenario
    month: int | None
    value_type: ValueType
    source_label: str
    component: str | None
    line_label: str | None          # internal only
    value: NormalizedValue
    validity: CellValidity
    external_dependency: bool
    title_conflict: bool            # the segment's title cell differs from the block title (historical columns)
    source_cell: str                # structural provenance (internal; never part of the key)
    cached_external_value: Any = None   # internal only: Excel's cached result of an externally dependent cell

    @property
    def usable(self) -> bool:
        return self.validity is CellValidity.VALID

    @property
    def provenance(self) -> str:
        return {CellKind.LITERAL: "literal", CellKind.FORMULA_VALUE: "cached_formula_result",
                CellKind.ERROR: "excel_error", CellKind.FORMULA_ERROR: "excel_error",
                CellKind.FORMULA_MISSING_CACHE: "formula_missing_cache"}.get(self.value.kind, "empty")


class RegionCoverage(BaseModel):
    model_config = ConfigDict(extra="forbid")
    normalized: bool
    cells: int = 0
    blank: int = 0
    valid: int = 0
    literal: int = 0
    cached_formula: int = 0
    excel_errors: dict[str, int] = Field(default_factory=dict)
    missing_cache: int = 0
    external_refs: int = 0
    type_mismatch: int = 0
    usable: int = 0
    invalid: int = 0


class BudgetSummary(BaseModel):
    model_config = ConfigDict(extra="forbid")
    schema_version: str | None = None
    planning_year: int | None = None
    planning_year_origin: str | None = None
    revision_label: str | None = None
    snapshot_date: dt.date | None = None
    snapshot_date_origin: str | None = None
    sheet: str | None = None
    sheets_excluded: list[str] = Field(default_factory=list)
    unexpected_sheets: int = 0
    bounds: str = BOUNDS
    blocks: list[str] = Field(default_factory=list)
    segments: int = 0
    historical_years: list[int] = Field(default_factory=list)
    expense_groups: int = 0
    summary_panels: int = 0
    cells_in_bounds: int = 0
    candidate_value_cells: int = 0
    blank_candidates: int = 0
    literal_numeric_cells: int = 0
    literal_text_cells: int = 0
    cached_formula_cells: int = 0
    formula_missing_cache: int = 0
    excel_errors: dict[str, int] = Field(default_factory=dict)
    ref_errors: int = 0
    div0_errors: int = 0
    other_errors: int = 0
    external_links: int = 0
    external_formula_cells: int = 0
    external_dependent_facts: int = 0
    external_dependency_regions: list[str] = Field(default_factory=list)
    usable_unverified: int = 0
    invalid_facts: int = 0
    invalid_by_validity: dict[str, int] = Field(default_factory=dict)
    invalid_cells_in_validated_regions: int = 0
    coverage_rate: float = 0.0     # usable / non-blank candidate value cells
    coverage_by_region: dict[str, RegionCoverage] = Field(default_factory=dict)
    validated_regions: dict[str, RegionCoverage] = Field(default_factory=dict)
    facts_by_region: dict[str, int] = Field(default_factory=dict)
    title_conflict_facts: int = 0
    top_header_conflicts: int = 0
    checks: dict[str, dict[str, int]] = Field(default_factory=dict)
    source_status: str | None = None
    contains_invalid_cells: bool = False
    authoritative: bool = False
    file_of_record: str = "unresolved"


@dataclass
class BudgetResult:
    ok: bool
    failed_step: str | None
    file: str
    facts: list[BudgetFact]
    report: ValidationReport
    verification: VerificationResult
    counts: BudgetSummary
    component: ImportResult | None

    @property
    def audience(self) -> Audience:
        return Audience.OWNER

    @property
    def usable_facts(self) -> list[BudgetFact]:
        return [f for f in self.facts if f.usable]

    @property
    def invalid_facts(self) -> list[BudgetFact]:
        return [f for f in self.facts if not f.usable]

    def summary(self) -> dict[str, Any]:
        """Structure and counts only. External link basenames are withheld (count only)."""
        comp = None
        if self.component is not None:
            comp = {"ok": self.component.ok, "failed_step": self.component.failed_step,
                    "report": _scrub(self.component.report.to_dict()),
                    "verification": self.component.verification.model_dump(mode="json")}
        return {"ok": self.ok, "failed_step": self.failed_step, "source": "butce", "audience": self.audience.value,
                "file": self.file, "counts": self.counts.model_dump(mode="json"), "report": _scrub(self.report.to_dict()),
                "component": comp, "verification": self.verification.model_dump(mode="json")}


def _scrub(report: dict[str, Any]) -> dict[str, Any]:
    links = report.get("external_links") or {}
    if links.get("names"):
        links["names"] = []
        links["names_withheld"] = True
    return report


# ---------------------------------------------------------------------------------------------
# File name / candidates
# ---------------------------------------------------------------------------------------------
def parse_filename(name: str) -> tuple[int, str] | None:
    m = BUTCE_FILENAME.match(unicodedata.normalize("NFC", name))
    return (int(m.group("year")), m.group("label")) if m else None


def find_butce_candidates(folder: str | Path) -> list[dict[str, Any]]:
    """List Bütçe-looking workbooks with schema compatibility and structural error counts. Never selects one:
    the file of record is an owner decision (no mtime, no name order, no error-count preference)."""
    out = []
    for p in sorted(Path(folder).glob("*.xlsx")):
        name = unicodedata.normalize("NFC", p.name)
        if name.startswith("~$") or "butce" not in name.lower():
            continue
        parsed = parse_filename(p.name)
        res = import_butce(p) if parsed else None
        c = res.counts if res else None
        out.append({"file": p.name, "planning_year": parsed[0] if parsed else None,
                    "revision_label": parsed[1] if parsed else None,
                    "schema_compatible": bool(res and res.ok), "failed_step": res.failed_step if res else "filename",
                    "source_status": c.source_status if c else None,
                    "excel_errors": dict(c.excel_errors) if c else {},
                    "external_formula_cells": c.external_formula_cells if c else 0,
                    "coverage_rate": c.coverage_rate if c else 0.0})
    return out


# ---------------------------------------------------------------------------------------------
# Pipeline
# ---------------------------------------------------------------------------------------------
def import_butce(path: str | Path, *, planning_year: int | None = None,
                 snapshot_date: dt.date | None = None) -> BudgetResult:
    """Import one explicitly chosen Bütçe workbook. Structural problems fail the import; invalid cells do not."""
    path = Path(path)
    report = ValidationReport(source="butce", audience=Audience.OWNER, file=path.name, layout="workbook")
    counts = BudgetSummary()
    result = BudgetResult(False, "year", path.name, [], report, VerificationResult(), counts, None)
    parsed = parse_filename(path.name)
    if parsed:
        counts.revision_label = parsed[1]
    if snapshot_date is not None:
        counts.snapshot_date, counts.snapshot_date_origin = snapshot_date, "declared"
    if planning_year is not None and parsed and parsed[0] != planning_year:
        report.add(Severity.ERROR, "planning_year_conflict", "The declared planning year differs from the filename year.")
        return result
    year = planning_year if planning_year is not None else (parsed[0] if parsed else None)
    if year is None:
        report.add(Severity.ERROR, "planning_year_unknown", "No declared planning year and no approved filename pattern.")
        return result
    counts.planning_year, counts.planning_year_origin = year, ("declared" if planning_year is not None else "filename")
    if year not in SUPPORTED_YEARS:
        report.add(Severity.ERROR, "unsupported_template", "No verified Bütçe template for this planning year.")
        return result
    counts.schema_version, counts.sheet = f"butce_{year}_v1", str(year)

    comp = run_import(path, butce_spec(year))
    result.component = comp
    counts.external_links = comp.report.external_links.count
    if not comp.ok:
        report.add(Severity.ERROR, "component_failed", "The Bütçe matrix failed structural validation.")
        result.failed_step = "component"
        return result
    _sheets(path, year, report, counts)

    grid: dict[str, tuple[str, NormalizedValue]] = {
        r.values["column_label"].value: (r.values["block"].value, r.values["value"]) for r in comp.records}
    group_titles = _check_layout(grid, year, report)
    if report.errors:
        result.failed_step = "layout"
        return result

    positions = _positions(year)
    facts = _facts(grid, positions)
    _coverage(grid, positions, facts, counts)
    _checks(grid, positions, facts, year, report, counts)
    _status(counts, report)
    counts.blocks = [b.block.value for b in BLOCKS]
    counts.segments = len(segments(year))
    counts.historical_years = sorted({s.year for s in segments(year) if s.year < year})
    counts.expense_groups = len(expense_groups(year))
    counts.summary_panels = len(PANELS)
    result.facts = facts
    _verify(result, grid, positions, group_titles, path)
    if not result.verification.ok:
        bad = [c.name for c in result.verification.checks if not c.ok]
        report.add(Severity.ERROR, "butce_verification_failed", "Verification failed: " + ", ".join(bad) + ".",
                   count=len(bad))
        result.failed_step, result.facts = "verify", []
        return result
    result.ok, result.failed_step = True, None
    return result


def _sheets(path: Path, year: int, report: ValidationReport, counts: BudgetSummary) -> None:
    try:
        with WorkbookPair(path) as wb:
            names = wb.sheet_names
    except Exception:  # run_import already validated readability; never leak a traceback
        return
    others = [n for n in names if n != str(year)]
    counts.sheets_excluded = [n for n in others if n in SCRATCH_SHEETS]
    counts.unexpected_sheets = len(others) - len(counts.sheets_excluded)
    if counts.unexpected_sheets:
        report.add(Severity.WARNING, "unexpected_sheets",
                   f"{counts.unexpected_sheets} sheet(s) besides the planning sheet and the scratch sheet; not read.",
                   count=counts.unexpected_sheets)


# ---------------------------------------------------------------------------------------------
# Layout (fatal on drift)
# ---------------------------------------------------------------------------------------------
def _text(grid, coord: str) -> str | None:
    entry = grid.get(coord)
    if entry is None:
        return None
    v = entry[1].value
    return v if isinstance(v, str) else None


def _nonblank(grid, coord: str) -> bool:
    entry = grid.get(coord)
    return entry is not None and (entry[1].value is not None or entry[1].error is not None)


def _check_row(grid, report: ValidationReport, row: int, c0: int, c1: int, expected: dict[int, str], what: str) -> None:
    """Every expected cell holds its exact label and every other cell in the row span is blank."""
    bad = sum(1 for c, lab in expected.items() if _text(grid, _coord(c, row)) != lab)
    extra = sum(1 for c in range(c0, c1 + 1) if c not in expected and _nonblank(grid, _coord(c, row)))
    if bad or extra:
        report.add(Severity.ERROR, "header_mismatch",
                   f"{what}: {bad} label cell(s) differ from the template and {extra} unexpected non-blank cell(s).",
                   region=f"{get_column_letter(c0)}{row}:{get_column_letter(c1)}{row}", count=bad + extra)


def _check_layout(grid, year: int, report: ValidationReport) -> dict[int, str]:
    segs = segments(year)
    a, bc = 1, _ci("BC")
    group_titles: dict[int, str] = {}
    known = {b.title for b in BLOCKS if b.title} | {HISTORICAL_TITLE}
    for b in BLOCKS:
        t = b.title_row
        title_expected = {_ci("B"): HISTORICAL_TITLE, _ci("Z"): HISTORICAL_TITLE}
        if b.title:
            title_expected.update({_ci(c): b.title for c in BLOCK_TITLE_COLS})
            _check_row(grid, report, t, a, bc, title_expected, f"block {b.index} title row")
        else:
            titles = [_text(grid, f"{c}{t}") for c in BLOCK_TITLE_COLS]
            ok = all(isinstance(x, str) and x.strip() for x in titles) and len(set(titles)) == 1 and titles[0] not in known
            if not ok:
                report.add(Severity.ERROR, "block_title_mismatch",
                           f"Block {b.index}: the group-list title cells are missing, differ from each other or repeat "
                           "another block's title.", region=f"AL{t}:AX{t}")
            else:
                group_titles[b.index] = titles[0]  # type: ignore[assignment]
            title_expected.update({_ci(c): titles[0] or "" for c in BLOCK_TITLE_COLS})
            _check_row(grid, report, t, a, bc, title_expected, f"block {b.index} title row")
        _check_row(grid, report, t + 1, a, bc, {s.col: s.label for s in segs}, f"block {b.index} segment row")
        units = {1: "Months"}
        for s in segs:
            u = segment_units(b, s)
            units.update({s.col: u[0], s.col + 1: u[1], s.col + 2: u[2]})
        _check_row(grid, report, t + 2, a, bc, units, f"block {b.index} unit row")
        labels = {t + 3 + i: m for i, m in enumerate(MONTHS_EN)} | {t + 15: "Total"}
        bad = sum(1 for r, lab in labels.items() if _text(grid, f"A{r}") != lab)
        if bad:
            report.add(Severity.ERROR, "row_label_mismatch", f"Block {b.index}: {bad} month/Total row label(s) differ.",
                       region=f"A{t + 3}:A{t + 15}", count=bad)
    if len(set(group_titles.values())) != len(group_titles):
        report.add(Severity.ERROR, "block_title_mismatch", "The two group-list blocks carry the same title.")

    groups = expense_groups(year)
    ed = _ci("ED")
    hdr, comp, unit = {}, {}, {}
    for j, (_, _, label) in enumerate(groups):
        g = EXPENSE_FIRST_COL + 4 * j
        hdr[g] = label
        for c, name in enumerate(EXPENSE_COMPONENTS):
            comp[g + c] = name
            unit[g + c] = "YTL"
    _check_row(grid, report, EXPENSE_HEADER_ROW, EXPENSE_FIRST_COL, ed, hdr, "expense group row")
    _check_row(grid, report, EXPENSE_COMPONENT_ROW, EXPENSE_FIRST_COL, ed, comp, "expense component row")
    _check_row(grid, report, EXPENSE_UNIT_ROW, EXPENSE_FIRST_COL, ed, unit, "expense unit row")
    bad = sum(1 for i, m in enumerate(MONTHS_EN) if _text(grid, f"A{EXPENSE_FIRST_MONTH_ROW + i}") != m)
    bad += _text(grid, f"A{EXPENSE_TOTAL_ROW}") != "Total"
    if bad:
        report.add(Severity.ERROR, "row_label_mismatch", f"{bad} expense month/Total row label(s) differ.",
                   region=f"A{EXPENSE_FIRST_MONTH_ROW}:A{EXPENSE_TOTAL_ROW}", count=bad)

    bad = 0
    for col, py, last, has_target in PANELS:
        c = _ci(col)
        expected = {(c, PANEL_TITLE_ROW): PANEL_TITLE, (c, EXPENSE_HEADER_ROW): f"{py} expenses"}
        expected |= {(c + 1 + k, PANEL_SCENARIO_ROW): lab for k, (lab, _) in enumerate(PANEL_SCENARIOS)}
        expected |= {(c, PANEL_FIRST_LINE_ROW + i): lab for i, lab in enumerate((*PANEL_LINES, last))}
        bad += sum(1 for (cc, r), lab in expected.items() if _text(grid, _coord(cc, r)) != lab)
        target = _text(grid, _coord(c, PANEL_TARGET_ROW))
        bad += (target != PANEL_TARGET_LABEL) if has_target else _nonblank(grid, _coord(c, PANEL_TARGET_ROW))
    if bad:
        report.add(Severity.ERROR, "panel_label_mismatch", f"{bad} BÜTÇE ÖZET panel label cell(s) differ from the template.",
                   region="CK24:ED31", count=bad)
    return group_titles


# ---------------------------------------------------------------------------------------------
# Facts
# ---------------------------------------------------------------------------------------------
def _positions(year: int) -> dict[str, _Pos]:
    pos: dict[str, _Pos] = {}
    segs = segments(year)
    for b in BLOCKS:
        t = b.title_row
        for s in segs:
            unit, rate1, rate2 = segment_units(b, s)
            block_title_cell = segment_title_cell(s) in BLOCK_TITLE_COLS
            conflict = not block_title_cell and b.block is not BudgetBlock.TOTAL_SALES
            for m in range(1, 13):
                pos[_coord(s.col, t + 2 + m)] = _Pos(BudgetRegion.SALES_MONTH, b.block, s.year, s.scenario, m,
                                                    ValueType.AMOUNT, unit, title_conflict=conflict)
            pos[_coord(s.col, t + 15)] = _Pos(BudgetRegion.SALES_TOTAL, b.block, s.year, s.scenario, None,
                                              ValueType.AMOUNT, unit, title_conflict=conflict)
            if b.index == 1:
                for k, lab in ((1, rate1), (2, rate2)):
                    for m in range(1, 13):
                        pos[_coord(s.col + k, t + 2 + m)] = _Pos(BudgetRegion.RATE_MONTH, None, s.year, s.scenario, m,
                                                                ValueType.CURRENCY_RATE, lab)
    for j, (gy, sc, _) in enumerate(expense_groups(year)):
        g = EXPENSE_FIRST_COL + 4 * j
        for c, comp in enumerate(EXPENSE_COMPONENTS):
            col, key = g + c, comp.lower()
            for m in range(1, 13):
                pos[_coord(col, EXPENSE_FIRST_MONTH_ROW + m - 1)] = _Pos(BudgetRegion.EXPENSE_MONTH, None, gy, sc, m,
                                                                         ValueType.AMOUNT, "YTL", key)
            pos[_coord(col, EXPENSE_TOTAL_ROW)] = _Pos(BudgetRegion.EXPENSE_TOTAL, None, gy, sc, None, ValueType.AMOUNT,
                                                       "YTL", key)
            pos[_coord(col, EXPENSE_INPUT_ROW)] = _Pos(BudgetRegion.EXPENSE_ANNUAL_INPUT, None, gy, sc, None,
                                                       ValueType.AMOUNT, "YTL", key)
            pos[_coord(col, EXPENSE_LINKED_ROW)] = _Pos(BudgetRegion.EXPENSE_LINKED_INPUT, None, gy, sc, None,
                                                        ValueType.AMOUNT, "YTL", key)
    for col, py, last, _ in PANELS:
        c = _ci(col)
        for k, (_, sc) in enumerate(PANEL_SCENARIOS):
            for i, line in enumerate((*PANEL_LINES, last)):
                pos[_coord(c + 1 + k, PANEL_FIRST_LINE_ROW + i)] = _Pos(BudgetRegion.BUDGET_SUMMARY, None, py, sc, None,
                                                                         ValueType.AMOUNT, PANEL_TITLE, line_label=line)
    return pos


def fact_key(p: _Pos) -> str:
    """Structural identity: region, block, year, scenario, month and the column/line identity. Never the file
    name, the revision label, the row number or Python `hash()`."""
    ident = p.source_label if p.region is BudgetRegion.RATE_MONTH else (p.component or p.line_label)
    return record_key("butce", [p.region.value, p.block.value if p.block else None, p.year, p.scenario.value, p.month,
                                ident])


def _contaminated_groups(grid, positions: dict[str, _Pos]) -> set[tuple[int, str]]:
    """Linked-input groups (row 1) that hold at least one external reference: the whole group is dependent."""
    out = set()
    for coord, p in positions.items():
        if p.region is BudgetRegion.EXPENSE_LINKED_INPUT and coord in grid and grid[coord][1].external_ref:
            out.add((p.year, p.scenario.value))
    return out


def _facts(grid, positions: dict[str, _Pos]) -> list[BudgetFact]:
    contaminated = _contaminated_groups(grid, positions)
    facts: list[BudgetFact] = []
    for coord, p in positions.items():
        nv = grid[coord][1]
        dependent = nv.external_ref or (p.region is BudgetRegion.EXPENSE_LINKED_INPUT
                                        and (p.year, p.scenario.value) in contaminated)
        if nv.value is None and nv.error is None and not nv.external_ref:
            continue  # blank: missing, never zero
        cached = None
        if nv.kind.is_error:
            validity = CellValidity.EXCEL_ERROR
        elif nv.error == "missing_cache":
            validity = CellValidity.MISSING_CACHE
        elif nv.error is not None:
            validity = CellValidity.TYPE_MISMATCH
        elif dependent:
            validity, cached = CellValidity.EXTERNAL_DEPENDENCY, nv.value
            nv = NormalizedValue(None, Confidence.INVALID, nv.kind, EXTERNAL_DEPENDENCY, nv.external_ref)
        elif isinstance(nv.value, bool) or not isinstance(nv.value, (int, float)):
            validity = CellValidity.TYPE_MISMATCH
            nv = NormalizedValue(None, Confidence.INVALID, nv.kind, TYPE_MISMATCH, nv.external_ref)
        else:
            validity = CellValidity.VALID
            if isinstance(nv.value, int):
                nv = NormalizedValue(float(nv.value), nv.confidence, nv.kind, None, nv.external_ref)
        if validity is not CellValidity.VALID and nv.confidence is not Confidence.INVALID:
            nv = NormalizedValue(None, Confidence.INVALID, nv.kind, nv.error, nv.external_ref)
        facts.append(BudgetFact(fact_key(p), p.region, p.block, p.year, p.scenario, p.month, p.value_type,
                                p.source_label, p.component, p.line_label, nv, validity, dependent, p.title_conflict,
                                coord, cached))
    return facts


# ---------------------------------------------------------------------------------------------
# Coverage
# ---------------------------------------------------------------------------------------------
def _coverage(grid, positions: dict[str, _Pos], facts: list[BudgetFact], counts: BudgetSummary) -> None:
    by_coord = {f.source_cell: f for f in facts}
    regions: dict[str, RegionCoverage] = {}
    validated: dict[str, RegionCoverage] = {}
    errors: Counter = Counter()
    for coord, (rng, nv) in grid.items():
        p = positions.get(coord)
        name = p.region.value if p else _validated_region(rng)
        target = regions if p else validated
        rc = target.setdefault(name, RegionCoverage(normalized=bool(p)))
        rc.cells += 1
        if nv.kind.is_formula and not nv.kind.is_error and nv.kind != CellKind.FORMULA_MISSING_CACHE:
            rc.cached_formula += 1
            counts.cached_formula_cells += 1
        elif nv.kind == CellKind.LITERAL:
            rc.literal += 1
            if isinstance(nv.value, str):
                counts.literal_text_cells += 1
            else:
                counts.literal_numeric_cells += 1
        if nv.external_ref:
            rc.external_refs += 1
            counts.external_formula_cells += 1
        if nv.kind.is_error:
            rc.excel_errors[nv.error] = rc.excel_errors.get(nv.error, 0) + 1
            errors[nv.error] += 1
        elif nv.kind == CellKind.FORMULA_MISSING_CACHE:
            rc.missing_cache += 1
            counts.formula_missing_cache += 1
        elif nv.error is not None:
            rc.type_mismatch += 1
        if p:
            f = by_coord.get(coord)
            if f is None:
                rc.blank += 1
            elif f.usable:
                rc.usable += 1
                rc.valid += 1
            else:
                rc.invalid += 1
                if f.validity is CellValidity.TYPE_MISMATCH and nv.error is None:
                    rc.type_mismatch += 1
        else:
            if nv.value is None and nv.error is None:
                rc.blank += 1
            elif nv.error is None and not nv.external_ref:
                rc.valid += 1
            else:
                rc.invalid += 1
    counts.coverage_by_region = dict(sorted(regions.items()))
    counts.validated_regions = dict(sorted(validated.items()))
    counts.cells_in_bounds = len(grid)
    counts.candidate_value_cells = len(positions)
    counts.blank_candidates = sum(r.blank for r in regions.values())
    counts.excel_errors = dict(sorted(errors.items()))
    counts.ref_errors = errors.get("#REF!", 0)
    counts.div0_errors = errors.get("#DIV/0!", 0)
    counts.other_errors = sum(errors.values()) - counts.ref_errors - counts.div0_errors
    counts.usable_unverified = sum(1 for f in facts if f.usable)
    counts.invalid_facts = sum(1 for f in facts if not f.usable)
    counts.invalid_by_validity = dict(sorted(Counter(f.validity.value for f in facts if not f.usable).items()))
    counts.invalid_cells_in_validated_regions = sum(r.invalid for r in validated.values())
    counts.external_dependent_facts = sum(1 for f in facts if f.external_dependency)
    counts.external_dependency_regions = sorted(
        {f.region.value for f in facts if f.external_dependency}
        | {name for name, r in validated.items() if r.external_refs})
    counts.facts_by_region = dict(sorted(Counter(f.region.value for f in facts).items()))
    counts.title_conflict_facts = sum(1 for f in facts if f.title_conflict)
    non_blank = counts.candidate_value_cells - counts.blank_candidates
    counts.coverage_rate = round(counts.usable_unverified / non_blank, 4) if non_blank else 0.0


# ---------------------------------------------------------------------------------------------
# Source checks (validation arithmetic only; never repairs; warnings / info)
# ---------------------------------------------------------------------------------------------
def _num(grid, coord: str) -> float | None:
    entry = grid.get(coord)
    if entry is None:
        return None
    nv = entry[1]
    if nv.error is not None or nv.external_ref or isinstance(nv.value, bool) or not isinstance(nv.value, (int, float)):
        return None
    return float(nv.value)


def _blank(grid, coord: str) -> bool:
    return not _nonblank(grid, coord)


def _close(a: float, b: float, tol: float = TOL) -> bool:
    return abs(round(a - b, 6)) <= tol


def _sum_check(grid, total: str, parts: list[str]) -> bool | None:
    """True/False when checkable; None when the total or a non-blank part is not a usable number."""
    t = _num(grid, total)
    if t is None:
        return None
    vals = []
    for p in parts:
        if _blank(grid, p):
            continue
        v = _num(grid, p)
        if v is None:
            return None
        vals.append(v)
    return _close(t, sum(vals))


def _checks(grid, positions, facts, year: int, report: ValidationReport, counts: BudgetSummary) -> None:
    segs = segments(year)
    checks: dict[str, Counter] = defaultdict(Counter)

    def tally(name: str, outcome: bool | None) -> None:
        if outcome is None:
            checks[name]["skipped"] += 1
        else:
            checks[name]["checked"] += 1
            checks[name]["failed"] += not outcome

    for b in BLOCKS:
        t = b.title_row
        for s in segs:
            tally("sales_total_vs_months", _sum_check(grid, _coord(s.col, t + 15),
                                                      [_coord(s.col, t + 2 + m) for m in range(1, 13)]))
    blk = {b.block: b.title_row for b in BLOCKS}
    roll = ((BudgetBlock.TOTAL_SALES, (BudgetBlock.TOTAL_SALES_WITHOUT_TROTEC, BudgetBlock.TROTEC)),
            (BudgetBlock.TOTAL_SALES_WITHOUT_TROTEC, (BudgetBlock.GROUP_LIST_1, BudgetBlock.GROUP_LIST_2,
                                                      BudgetBlock.TRODAT, BudgetBlock.IMALAT)))
    for s in segs:
        for m in range(1, 13):
            for whole, parts in roll:
                coords = [_coord(s.col, blk[x] + 2 + m) for x in (whole, *parts)]
                vals = [_num(grid, c) for c in coords]
                tally("block_rollup", None if any(v is None for v in vals) else _close(vals[0], sum(vals[1:])))
            for b in BLOCKS[1:]:
                for k in (1, 2):
                    base = _num(grid, _coord(s.col + k, 4 + 2 + m))
                    copy = _num(grid, _coord(s.col + k, b.title_row + 2 + m))
                    tally("group_rate_columns_match_block_1",
                          None if base is None or copy is None else _close(base, copy, RATE_TOL))
    for j, _ in enumerate(expense_groups(year)):
        g = EXPENSE_FIRST_COL + 4 * j
        for c in range(3):
            tally("expense_total_vs_months", _sum_check(grid, _coord(g + c, EXPENSE_TOTAL_ROW),
                                                        [_coord(g + c, r) for r in range(7, 19)]))
        for r in range(EXPENSE_FIRST_MONTH_ROW, EXPENSE_TOTAL_ROW + 1):
            vals = [_num(grid, _coord(g + c, r)) for c in range(3)]
            if all(_blank(grid, _coord(g + c, r)) for c in range(3)):
                continue
            tally("expense_total_component", None if any(v is None for v in vals) else _close(vals[2], vals[0] + vals[1]))
    # Rows 1–2 repeat the segment/unit header above block 1; count disagreements (the block rows are used).
    top = sum(1 for s in segs if _text(grid, _coord(s.col, 1)) != _text(grid, _coord(s.col, 5)))
    top += sum(1 for c in range(2, _ci("BC") + 1) if _text(grid, _coord(c, 2)) != _text(grid, _coord(c, 6)))
    counts.top_header_conflicts = top
    counts.checks = {k: {"checked": v["checked"], "failed": v["failed"], "skipped": v["skipped"]}
                     for k, v in sorted(checks.items())}

    texts = {
        "sales_total_vs_months": "sales Total-row cell(s) differ from the sum of their months",
        "block_rollup": "block roll-up check(s) differ (total = without-Trotec + Trotec; without-Trotec = four groups)",
        "group_rate_columns_match_block_1": "rate cell(s) in blocks 2–7 differ from block 1",
        "expense_total_vs_months": "expense Total-row cell(s) differ from the sum of their months",
        "expense_total_component": "expense Total component(s) differ from Finance + Other",
    }
    for name, text in texts.items():
        n = checks[name]["failed"]
        if n:
            sev = Severity.INFO if name == "group_rate_columns_match_block_1" else Severity.WARNING
            report.add(sev, f"check_{name}", f"{n} {text}; values are kept as recorded.", count=n)
    if top:
        report.add(Severity.INFO, "top_header_conflicts",
                   f"{top} cell(s) of the duplicate header in rows 1–2 differ from block 1's header; identity uses each "
                   "block's own header rows.", count=top)
    if counts.title_conflict_facts:
        report.add(Severity.INFO, "historical_title_conflict",
                   f"{counts.title_conflict_facts} fact(s) sit in historical columns titled '{HISTORICAL_TITLE}' inside a "
                   "block with another title; semantics unverified, kept with a flag.", count=counts.title_conflict_facts)


def _status(counts: BudgetSummary, report: ValidationReport) -> None:
    invalid_anywhere = counts.invalid_facts + counts.invalid_cells_in_validated_regions + counts.external_formula_cells
    counts.contains_invalid_cells = invalid_anywhere > 0
    if counts.usable_unverified == 0:
        counts.source_status = SourceStatus.UNUSABLE.value
    elif invalid_anywhere:
        counts.source_status = SourceStatus.PARTIALLY_USABLE.value
    else:
        counts.source_status = SourceStatus.COMPLETE.value
    if counts.invalid_facts:
        report.add(Severity.WARNING, "invalid_facts",
                   f"{counts.invalid_facts} candidate value cell(s) are invalid (Excel error, missing cache, external "
                   "dependency or non-numeric) and excluded from usable coverage.", count=counts.invalid_facts)
    if counts.invalid_cells_in_validated_regions:
        report.add(Severity.WARNING, "invalid_cells_in_validated_regions",
                   f"{counts.invalid_cells_in_validated_regions} invalid cell(s) in validated-only regions (derived rows, "
                   "deviation panels, side cells).", count=counts.invalid_cells_in_validated_regions)
    if counts.external_dependent_facts or counts.external_formula_cells:
        report.add(Severity.WARNING, "external_dependency",
                   f"{counts.external_formula_cells} formula cell(s) reference another workbook; "
                   f"{counts.external_dependent_facts} fact(s) depend on it and are unusable. The linked workbook is not "
                   "opened.", count=max(counts.external_formula_cells, counts.external_dependent_facts))
    report.add(Severity.INFO, "source_not_authoritative",
               f"Source status {counts.source_status}; the file of record is unresolved (owner decision).")


# ---------------------------------------------------------------------------------------------
# Verification
# ---------------------------------------------------------------------------------------------
def _area(ref: str) -> int:
    c0, r0, c1, r1 = range_boundaries(ref)
    return (c1 - c0 + 1) * (r1 - r0 + 1)


def _verify(result: BudgetResult, grid, positions, group_titles: dict[int, str], path: Path) -> None:
    v, facts, counts = result.verification, result.facts, result.counts
    comp = result.component
    v.add("audience_owner", result.audience is Audience.OWNER and comp is not None and comp.audience is Audience.OWNER)
    keys = [f.record_key for f in facts]
    v.add("fact_keys_unique", len(keys) == len(set(keys)), f"{len(set(keys))} distinct")
    year = counts.planning_year or 0
    area = _area(BOUNDS)
    v.add("region_map_complete", len(grid) == area and sum(_area(r.range) for r in butce_ranges(year)) == area,
          f"{len(grid)} of {area}")
    blank = sum(r.blank for r in counts.coverage_by_region.values())
    v.add("coverage_arithmetic", counts.usable_unverified + counts.invalid_facts + blank == counts.candidate_value_cells)
    v.add("invalid_never_usable", all((f.validity is CellValidity.VALID) == f.usable for f in facts)
          and all(f.value.value is None and f.value.confidence is Confidence.INVALID for f in facts if not f.usable))
    v.add("values_never_trusted", all(f.value.confidence in (Confidence.UNVERIFIED, Confidence.INVALID) for f in facts))
    v.add("external_never_usable", not any(f.external_dependency and f.usable for f in facts))
    err_cells = sum(1 for c in positions if grid[c][1].kind.is_error)
    v.add("errors_preserved", err_cells == sum(1 for f in facts if f.validity is CellValidity.EXCEL_ERROR),
          f"{err_cells} error cell(s)")
    blob = json.dumps(result.summary(), ensure_ascii=False, default=str)
    secrets: set[str] = set(group_titles.values())
    for f in facts:
        for val in (f.value.value, f.cached_external_value):
            if isinstance(val, (int, float)) and not isinstance(val, bool) and (
                    abs(val) >= 1000 or (abs(val) >= 1 and not float(val).is_integer())):
                secrets.add(repr(float(val)))
        if isinstance(f.cached_external_value, str) and len(f.cached_external_value) >= 4:
            secrets.add(f.cached_external_value)
    for name in (comp.report.external_links.names if comp else []):
        secrets.add(name)
    parent = str(path.resolve().parent)
    if len(parent) > 1:
        secrets.add(parent)
    leaked = sum(1 for s in secrets if s and s in blob)
    v.add("summary_has_no_values_labels_or_paths", leaked == 0, f"{leaked} leak(s)")


__all__ = ["BudgetBlock", "BudgetFact", "BudgetRegion", "BudgetResult", "BudgetSummary", "CellValidity",
           "RegionCoverage", "Scenario", "SourceStatus", "ValueType", "BLOCKS", "PANELS", "SUPPORTED_YEARS",
           "butce_ranges", "butce_spec", "expense_groups", "fact_key", "find_butce_candidates", "import_butce",
           "parse_filename", "segments", "segment_units"]
